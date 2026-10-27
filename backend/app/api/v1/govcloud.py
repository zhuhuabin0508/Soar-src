"""政务云资产 API：按资源类型分表，Excel「原始数据」列一比一入库。

不导入 Sheet「数据统计」及统计类表头；分布统计由库内字段自行聚合（前端暂不展示）。
"""
from __future__ import annotations

import csv
import io
import json
import logging
import time
from datetime import datetime
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import func, or_, text
from sqlalchemy.dialects.postgresql import insert
from sqlalchemy.orm import Session

from app.core.audit import get_client_ip, log_audit
from app.core.govcloud_schema import (
    HEADER_ALIAS,
    EXTRA_IP_KEYS,
    EXTRA_NAME_KEYS,
    FIELD_LABELS,
    RESERVED_COLUMNS,
    RESOURCE_EXCEL_FIELDS,
    SOURCE_ID_KEYS,
    excel_attr,
    norm_header_key,
    skip_excel_header,
)
from app.core.qingteng_asset_sync import get_qingteng_devices, sync_device_hosts
from app.core.timezone import beijing_now
from app.database import get_db
from app.dependencies import require_permission
from app.models.device import Device
from app.models.govcloud_asset import GovCloudImportBatch, RESOURCE_MODELS
from app.models.qingteng_host import (
    QingtengHostAsset,
    QINGTENG_FIELD_LABELS,
    QINGTENG_HOST_FIELDS,
    QINGTENG_LIST_FIELDS,
)
from app.models.user import User

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/govcloud", tags=["govcloud"])

RESOURCE_TYPES = {
    "cloud_host": {
        "label": "云主机",
        "excel_name": "云主机资产",
        "source_group": "itsm_ops",
        "filename_hints": ("云主机", "cloud_host"),
        "search_fields": (
            "instance_name", "inst_name", "ip_address", "floatip_address",
            "user_department", "applicant_department", "source_id",
        ),
    },
    "bare_metal": {
        "label": "裸金属",
        "excel_name": "裸金属资产",
        "source_group": "itsm_ops",
        "filename_hints": ("裸金属", "bare_metal"),
        "search_fields": (
            "instance_name", "inst_name", "ip_address", "floatip_address",
            "user_department", "applicant_department", "source_id",
        ),
    },
    "e_government_network": {
        "label": "网络资源",
        "excel_name": "网络资源资产",
        "source_group": "itsm_ops",
        "filename_hints": ("网络资源", "e_government_network", "network"),
        "search_fields": (
            "instance_name", "inst_name", "name", "ip_address",
            "user_department", "applicant_department", "source_id",
        ),
    },
    "elastic_ip": {
        "label": "弹性IP",
        "excel_name": "弹性IP资产",
        "source_group": "itsm_ops",
        "filename_hints": ("弹性IP", "弹性ip", "elastic_ip"),
        "search_fields": (
            "instance_name", "inst_name", "ip_address", "floatip_address",
            "user_department", "applicant_department", "source_id",
        ),
    },
    # 青藤主机：凭据来自设备管理（auth_type=qingteng），API 拉取入库而非 Excel 导入。
    # resource_type 实际取值带设备维度："qingteng_host:<设备ID>"（每台青藤一张数据表卡片）。
    "qingteng_host": {
        "label": "青藤主机",
        "excel_name": "青藤主机资产",
        "source_group": "qingteng_host",
        "filename_hints": (),
        "search_fields": (
            "hostname", "display_ip", "connection_ip", "external_ip",
            "agent_id", "biz_group", "serial_number", "charge_name", "remark",
        ),
    },
}

# 青藤主机模型注册（数据表卡片由设备列表动态派生，见 meta/summary）
RESOURCE_MODELS["qingteng_host"] = QingtengHostAsset

SOURCE_GROUPS = {
    "itsm_ops": {
        "label": "综合运管平台资产发现",
        "index": "",
        "enabled": True,
        "sort_order": 1,
        "description": "综合运管平台每周全量爬取，对应云主机 / 裸金属 / 网络资源 / 弹性IP 四张 Excel",
    },
    "intranet_mapping": {
        "label": "内网资产测绘",
        "index": "1.1",
        "enabled": False,
        "sort_order": 2,
        "description": "内网资产测绘结果，后续接入",
    },
    "qingteng_host": {
        "label": "青藤主机资产发现",
        "index": "1.2",
        "enabled": False,  # 实际启用状态按设备列表动态派生（meta/summary 内覆盖）
        "sort_order": 3,
        "description": "登录设备管理中配置的青藤平台（支持多台），API 拉取 linux / win 主机清单入库",
    },
    "jiaotu_yunsuo": {
        "label": "椒图云锁主机资产发现",
        "index": "1.3",
        "enabled": False,
        "sort_order": 4,
        "description": "椒图云锁主机资产发现结果，后续接入",
    },
    "tianqing_endpoint": {
        "label": "天擎终端资产发现",
        "index": "1.4",
        "enabled": False,
        "sort_order": 5,
        "description": "天擎终端资产发现结果，后续接入",
    },
}

MAX_ROWS = 250000
MAX_FILE_SIZE = 200 * 1024 * 1024
UPSERT_CHUNK = 500
# 导出流式分块行数：每块一次查询，避免一次性加载全表
EXPORT_CHUNK = 2000
# 导出尾部追加的业务列（与列表页展示口径一致，便于追溯入库/同步时间）
EXPORT_TAIL_COLUMNS = ("source_id", "created_at", "updated_at")
# 青藤导出尾部业务列（额外带来源青藤平台名，多平台数据可追溯）
QT_EXPORT_TAIL_COLUMNS = ("source_id", "qingteng_device_name", "created_at", "updated_at")


def _keyword_clauses(model, meta: dict, kw: str) -> list:
    """构造 keyword 模糊搜索条件（列表与导出共用同一口径）。"""
    like = f"%{kw}%"
    clauses = [model.source_id.ilike(like)]
    if hasattr(model, "cmdb_id"):
        clauses.append(model.cmdb_id.ilike(like))
    for field in meta.get("search_fields") or []:
        if field in ("source_id", "_id"):
            continue
        col = getattr(model, excel_attr(field), None)
        if col is not None:
            clauses.append(col.ilike(like))
    extra = getattr(model, "extra_excel", None)
    if extra is not None:
        for key in (*EXTRA_IP_KEYS, *EXTRA_NAME_KEYS):
            clauses.append(extra[key].astext.ilike(like))
    return clauses


def _csv_val(val) -> str:
    """导出单元格序列化：None→空串，datetime→YYYY-mm-dd HH:MM:SS。"""
    if val is None:
        return ""
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d %H:%M:%S")
    return str(val)


def _row_cells(
    row,
    fields: list[str],
    extra_keys: list[str],
    extra_col: str = "extra_excel",
    tail: tuple = EXPORT_TAIL_COLUMNS,
) -> list[str]:
    """导出行序列化：字段列 + extra 追加列 + 业务尾部列（source_id/时间等）。"""
    cells = [_csv_val(getattr(row, excel_attr(f), "")) for f in fields]
    extra = getattr(row, extra_col, None)
    extra = extra if isinstance(extra, dict) else {}
    cells.extend(_csv_val(extra.get(k)) for k in extra_keys)
    cells.extend(_csv_val(getattr(row, k, "")) for k in tail)
    return cells


def _resolve_type(resource_type: str) -> tuple[str, dict, int | None]:
    """解析资源类型：支持青藤的设备维度 ``qingteng_host:<设备ID>``。

    返回 (基础类型, 类型元数据, 设备 ID)。青藤之外的多余后缀视为非法。
    """
    base, _, dev = (resource_type or "").partition(":")
    meta = RESOURCE_TYPES.get(base)
    if not meta:
        raise HTTPException(status_code=400, detail="不支持的资源类型")
    if ":" in (resource_type or ""):
        if base != "qingteng_host":
            raise HTTPException(status_code=400, detail="该资源类型不支持设备维度参数")
        if not dev or not dev.isdigit():
            raise HTTPException(status_code=400, detail="青藤资源类型格式应为 qingteng_host:<设备ID>")
        return base, meta, int(dev)
    return base, meta, None


def _require_type(resource_type: str) -> dict:
    """按基础类型校验（兼容旧调用；带设备后缀时取基础类型元数据）。"""
    base, meta, _ = _resolve_type(resource_type)
    return meta


def _model(resource_type: str):
    base, _, _ = _resolve_type(resource_type)
    return RESOURCE_MODELS[base]


def _device_filter(model, base: str, device_id: int | None):
    """青藤模型按设备维度过滤（None 表示全部设备）。其它类型返回 None。"""
    if base == "qingteng_host" and device_id is not None:
        return model.qingteng_device_id == device_id
    return None


def _qt_columns() -> list[dict]:
    """青藤主机列表页展示列（详情/导出用全量字段，见 export）。"""
    return [{"key": k, "label": QINGTENG_FIELD_LABELS.get(k, k)} for k in QINGTENG_LIST_FIELDS]


def _columns(resource_type: str) -> list[dict]:
    base, _, _ = _resolve_type(resource_type)
    if base == "qingteng_host":
        return _qt_columns()
    cols = []
    for key in RESOURCE_EXCEL_FIELDS[base]:
        cols.append({"key": key, "label": FIELD_LABELS.get(key, key)})
    return cols


def _cell_value(val) -> str:
    if val is None:
        return ""
    if isinstance(val, (dict, list)):
        return json.dumps(val, ensure_ascii=False)
    if isinstance(val, bool):
        return "true" if val else "false"
    return str(val).strip()


def _normalize_header(key: str) -> str:
    """表头归一化：中文标签（如导出 CSV 的表头「实例名称/IP」）与英文字段名的
    空格/大小写/下划线变体（「Instance Name」「INSTANCE_NAME」）映射回英文字段名，
    使导出口径的文件可直接再导入；未知表头原样返回（进 extra_excel），
    统计类表头不受影响。键口径与 HEADER_ALIAS 构建共用 norm_header_key。
    """
    k = (key or "").strip()
    if not k or skip_excel_header(k):
        return k
    field = HEADER_ALIAS.get(norm_header_key(k))
    return field if field else k


def _decode_csv_bytes(raw: bytes) -> str:
    """CSV 解码：UTF-8（含 BOM）优先，GBK/GB18030 兜底（Excel 中文另存常见）。"""
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass
    try:
        return raw.decode("gb18030")
    except UnicodeDecodeError:
        raise HTTPException(
            status_code=400,
            detail="CSV 编码无法解析（支持 UTF-8 / GBK，请勿使用 UTF-16 等格式另存）",
        )


def _rows_from_grid(rows_iter) -> list[dict]:
    """CSV / Excel 共用的行构建：首行表头归一化，后续行按键取值。

    口径：短行补空、空行跳过、统计类表头列忽略、超过 MAX_ROWS 拒绝，
    两种格式走同一实现，避免口径漂移。
    """
    try:
        headers = [_normalize_header(_cell_value(h)) for h in next(rows_iter)]
    except StopIteration:
        return []
    result: list[dict] = []
    for row in rows_iter:
        item = {}
        empty = True
        for i, key in enumerate(headers):
            if skip_excel_header(key):
                continue
            val = _cell_value(row[i] if i < len(row) else "")
            if val:
                empty = False
            item[key] = val
        if not empty:
            result.append(item)
        if len(result) > MAX_ROWS:
            raise HTTPException(status_code=413, detail=f"行数过多，最多 {MAX_ROWS} 行")
    return result


def _parse_csv(raw: bytes) -> list[dict]:
    """解析 CSV（首行表头），行结构与 ``_parse_excel`` 同一套口径。"""
    return _rows_from_grid(csv.reader(io.StringIO(_decode_csv_bytes(raw))))


def _parse_excel(raw: bytes) -> list[dict]:
    """只读 Sheet「原始数据」，跳过「数据统计」页和统计类表头。"""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = None
    for sheet in wb.worksheets:
        title = (sheet.title or "").strip()
        if title == "数据统计":
            continue
        if title == "原始数据":
            ws = sheet
            break
    if ws is None:
        for sheet in wb.worksheets:
            if (sheet.title or "").strip() != "数据统计":
                ws = sheet
                break
    if ws is None:
        return []
    return _rows_from_grid(ws.iter_rows(values_only=True))


def _source_id(row: dict) -> str:
    for key in SOURCE_ID_KEYS:
        val = _cell_value(row.get(key))
        if val:
            return val[:128]
    ip = _cell_value(row.get("ip_address") or row.get("ip"))
    name = _cell_value(row.get("instance_name") or row.get("inst_name") or row.get("name"))
    if ip or name:
        return f"{ip}|{name}"[:128]
    return ""


def _row_to_values(resource_type: str, row: dict, batch_id: int, now) -> dict:
    fields = RESOURCE_EXCEL_FIELDS[resource_type]
    known = set(fields)
    values = {
        "source_id": _source_id(row),
        "batch_id": batch_id,
        "created_at": now,
        "updated_at": now,
    }
    extra = {}
    for key, raw in row.items():
        if skip_excel_header(key):
            continue
        val = _cell_value(raw)
        if key in RESERVED_COLUMNS:
            if key == "source_id":
                # 唯一标识已由 _source_id 使用（导出 CSV 回导场景），不再进 extra
                continue
            extra[key] = val
            continue
        if key in known:
            values[excel_attr(key)] = val
        else:
            extra[key] = val
    for field in fields:
        attr = excel_attr(field)
        values.setdefault(attr, "")
    values["extra_excel"] = extra
    return values


def infer_resource_type(filename: str, explicit: str) -> str:
    if explicit:
        _require_type(explicit)
        return explicit
    name = (filename or "").lower()
    for code, meta in RESOURCE_TYPES.items():
        for hint in meta["filename_hints"]:
            if hint.lower() in name:
                return code
    raise HTTPException(
        status_code=400,
        detail="无法从文件名识别资源类型，请选择：云主机 / 裸金属 / 网络资源 / 弹性IP",
    )


def _excel_columns(resource_type: str) -> list[dict]:
    cols = []
    base, _, _ = _resolve_type(resource_type)
    for key in RESOURCE_EXCEL_FIELDS[base]:
        cols.append({"key": key, "label": FIELD_LABELS.get(key, key)})
    return cols


def _type_payload(code: str, meta: dict, device: Device | None = None) -> dict:
    """资源类型 payload；青藤按设备生成（code 带设备维度，卡片名 = 设备名）。"""
    real_code = f"{code}:{device.id}" if device else code
    return {
        "code": real_code,
        "label": (device.name if device else meta["label"]),
        "excel_name": (device.name if device else meta.get("excel_name")) or meta["label"],
        "source_group": meta.get("source_group", "itsm_ops"),
        "excel_columns": _columns(real_code),
        "device_id": device.id if device else None,
    }


def _qingteng_table_payloads(db: Session) -> list[dict]:
    """青藤数据表卡片：每台已启用的青藤设备一张（多青藤零配置纳管）。"""
    return [
        _type_payload("qingteng_host", RESOURCE_TYPES["qingteng_host"], device=dev)
        for dev in get_qingteng_devices(db)
    ]


@router.get("/meta")
def govcloud_meta(
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("asset_discovery", "view")),  # noqa: ARG001
) -> dict:
    qingteng_tables = _qingteng_table_payloads(db)
    types = [
        _type_payload(code, meta)
        for code, meta in RESOURCE_TYPES.items()
        if code != "qingteng_host"
    ]
    types.extend(qingteng_tables)
    groups = []
    for gcode, gmeta in SOURCE_GROUPS.items():
        if gcode == "qingteng_host":
            tables, enabled = qingteng_tables, bool(qingteng_tables)
        else:
            tables = [t for t in types if t["source_group"] == gcode]
            enabled = bool(gmeta.get("enabled", True))
        groups.append({
            "code": gcode,
            "label": gmeta["label"],
            "index": gmeta.get("index", ""),
            "enabled": enabled,
            "sort_order": gmeta.get("sort_order", 0),
            "description": gmeta.get("description", ""),
            "tables": tables,
        })
    groups.sort(key=lambda g: g["sort_order"])
    return {
        "types": types,
        "groups": groups,
        "field_labels": {**FIELD_LABELS, **QINGTENG_FIELD_LABELS},
    }


@router.get("/summary")
def govcloud_summary(
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("asset_discovery", "view")),  # noqa: ARG001
) -> dict:
    groups = []
    qingteng_devices = get_qingteng_devices(db)
    for gcode, gmeta in SOURCE_GROUPS.items():
        tables = []
        if gcode == "qingteng_host":
            # 每台青藤设备一张数据表卡片：总数/一致性按设备维度统计
            for dev in qingteng_devices:
                code = f"qingteng_host:{dev.id}"
                live = int(
                    db.query(func.count(QingtengHostAsset.id))
                    .filter(QingtengHostAsset.qingteng_device_id == dev.id)
                    .scalar() or 0
                )
                last_batch = (
                    db.query(GovCloudImportBatch)
                    .filter(GovCloudImportBatch.resource_type == code)
                    .order_by(GovCloudImportBatch.id.desc())
                    .first()
                )
                unique_ids = last_batch.unique_source_ids if last_batch else 0
                consistent = (live == unique_ids and (last_batch.skipped == 0 if last_batch else True))
                if last_batch is None:
                    consistent = live == 0
                tables.append({
                    "code": code,
                    "label": dev.name,
                    "excel_name": dev.name,
                    "excel_columns": _qt_columns(),
                    "total": live,
                    "consistent": consistent,
                    "last_batch": last_batch.to_dict() if last_batch else None,
                    "device_id": dev.id,
                })
            groups.append({
                "code": gcode,
                "label": gmeta["label"],
                "index": gmeta.get("index", ""),
                "enabled": bool(qingteng_devices),
                "sort_order": gmeta.get("sort_order", 0),
                "description": gmeta.get("description", ""),
                "tables": tables,
            })
            continue
        for code, meta in RESOURCE_TYPES.items():
            if meta.get("source_group") != gcode or code == "qingteng_host":
                continue
            model = RESOURCE_MODELS[code]
            live = int(db.query(func.count(model.id)).scalar() or 0)
            last_batch = (
                db.query(GovCloudImportBatch)
                .filter(GovCloudImportBatch.resource_type == code)
                .order_by(GovCloudImportBatch.id.desc())
                .first()
            )
            unique_ids = last_batch.unique_source_ids if last_batch else 0
            consistent = (live == unique_ids and (last_batch.skipped == 0 if last_batch else True))
            if last_batch is None:
                consistent = live == 0
            tables.append({
                "code": code,
                "label": meta["label"],
                "excel_name": meta.get("excel_name") or meta["label"],
                "excel_columns": _columns(code),
                "total": live,
                "consistent": consistent,
                "last_batch": last_batch.to_dict() if last_batch else None,
            })
        groups.append({
            "code": gcode,
            "label": gmeta["label"],
            "index": gmeta.get("index", ""),
            "enabled": bool(gmeta.get("enabled", True)),
            "sort_order": gmeta.get("sort_order", 0),
            "description": gmeta.get("description", ""),
            "tables": tables,
        })
    groups.sort(key=lambda g: g["sort_order"])
    return {"groups": groups}


@router.get("/assets")
def list_govcloud_assets(
    resource_type: str = Query(..., description="资源类型（青藤为 qingteng_host:<设备ID>）"),
    keyword: str = Query("", description="名称/IP/部门/source_id 模糊搜索"),
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=200),
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("asset_discovery", "view")),  # noqa: ARG001
) -> dict:
    base, meta, device_id = _resolve_type(resource_type)
    model = _model(resource_type)
    query = db.query(model)
    dev_filter = _device_filter(model, base, device_id)
    if dev_filter is not None:
        query = query.filter(dev_filter)
    kw = keyword.strip()
    if kw:
        query = query.filter(or_(*_keyword_clauses(model, meta, kw)))

    total = query.count()
    rows = (
        query.order_by(model.id.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
        .all()
    )
    return {
        "total": total,
        "page": page,
        "page_size": page_size,
        "columns": _columns(resource_type),
        "items": [r.to_dict() for r in rows],
    }


@router.get("/export")
def export_govcloud_assets(
    resource_type: str = Query(..., description="资源类型（青藤为 qingteng_host:<设备ID>）"),
    keyword: str = Query("", description="与列表一致的搜索条件（可选）"),
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("asset_discovery", "view")),
    request: Request = None,
) -> StreamingResponse:
    """导出指定资源类型为 CSV（UTF-8 BOM，Excel 可直接打开）。

    - 列：该类型 Excel 字段（field_label 作表头）+ extra_excel 出现过的
      额外列（追加尾部，爬虫新增列不丢）+ source_id/入库时间/更新时间；
      青藤为全量主机字段 + extra + 来源青藤平台
    - 行：按 id 升序全量导出；传 keyword 时与列表搜索口径一致
    - 流式分块写出，避免一次性加载全表（单次导入上限 25 万行）
    """
    base, meta, device_id = _resolve_type(resource_type)
    model = _model(resource_type)
    kw = keyword.strip()
    kw_clauses = _keyword_clauses(model, meta, kw) if kw else []
    dev_filter = _device_filter(model, base, device_id)

    if base == "qingteng_host":
        fields = list(QINGTENG_HOST_FIELDS)
        labels = QINGTENG_FIELD_LABELS
        extra_col = "extra"
        tail = QT_EXPORT_TAIL_COLUMNS
    else:
        fields = RESOURCE_EXCEL_FIELDS[base]
        labels = FIELD_LABELS
        extra_col = "extra_excel"
        tail = EXPORT_TAIL_COLUMNS
    # extra 中出现过的额外列（jsonb_object_keys 汇总去重），保证新列不丢；
    # 排除与尾部业务列同名的键，避免 CSV 出现重复表头
    extra_keys = sorted(
        k for k in db.execute(
            text(f"SELECT DISTINCT k FROM {model.__tablename__}, jsonb_object_keys({extra_col}) AS k")
        ).scalars()
        if k not in fields and k not in tail
    )
    headers = (
        [labels.get(f, f) for f in fields]
        + [labels.get(k, k) for k in extra_keys]
        + list(tail)
    )

    # 审计：导出操作（GET 请求不被中间件捕获，需手动记录）
    try:
        log_audit(
            db,
            user_id=user.id,
            username=user.username,
            action="export",
            resource_type="asset_discovery",
            detail=f"{resource_type} keyword={kw!r}",
            ip_address=get_client_ip(request) if request else None,
            result="success",
        )
    except Exception as exc:  # noqa: BLE001
        logger.warning("导出审计写入失败（忽略）: %s", exc)

    def _iter_csv():
        buf = io.StringIO()
        buf.write("\ufeff")  # BOM：Excel 直接打开不乱码
        csv.writer(buf).writerow(headers)
        yield buf.getvalue()

        last_id = 0
        exported = 0
        while True:
            query = db.query(model).filter(model.id > last_id)
            if kw_clauses:
                query = query.filter(or_(*kw_clauses))
            if dev_filter is not None:
                query = query.filter(dev_filter)
            rows = query.order_by(model.id).limit(EXPORT_CHUNK).all()
            if not rows:
                break
            chunk = io.StringIO()
            writer = csv.writer(chunk)
            for row in rows:
                last_id = row.id
                exported += 1
                writer.writerow(_row_cells(row, fields, extra_keys, extra_col, tail))
            yield chunk.getvalue()
            if len(rows) < EXPORT_CHUNK:
                break
        logger.info("资产发现导出: type=%s keyword=%r rows=%d", resource_type, kw, exported)

    excel_name = meta.get("excel_name") or meta["label"]
    filename = f"{excel_name}_{beijing_now().strftime('%Y%m%d_%H%M%S')}.csv"
    return StreamingResponse(
        _iter_csv(),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@router.get("/assets/{resource_type}/{asset_id}")
def get_govcloud_asset(
    resource_type: str,
    asset_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("asset_discovery", "view")),  # noqa: ARG001
) -> dict:
    model = _model(resource_type)
    rec = db.query(model).filter(model.id == asset_id).first()
    if rec is None:
        raise HTTPException(status_code=404, detail="记录不存在")
    return rec.to_dict()


@router.get("/batches")
def list_batches(
    resource_type: str = Query("", description="可选，按类型筛选"),
    limit: int = Query(20, ge=1, le=100),
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("asset_discovery", "view")),  # noqa: ARG001
) -> dict:
    query = db.query(GovCloudImportBatch)
    if resource_type:
        _require_type(resource_type)
        query = query.filter(GovCloudImportBatch.resource_type == resource_type)
    rows = query.order_by(GovCloudImportBatch.id.desc()).limit(limit).all()
    return {"items": [r.to_dict() for r in rows]}


@router.get("/consistency")
def check_consistency(
    resource_type: str = Query(..., description="资源类型（青藤为 qingteng_host:<设备ID>）"),
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("asset_discovery", "view")),  # noqa: ARG001
) -> dict:
    base, _, device_id = _resolve_type(resource_type)
    model = _model(resource_type)
    dev_filter = _device_filter(model, base, device_id)
    live_query = db.query(func.count(model.id))
    if dev_filter is not None:
        live_query = live_query.filter(dev_filter)
    live = int(live_query.scalar() or 0)
    last_full = (
        db.query(GovCloudImportBatch)
        .filter(
            GovCloudImportBatch.resource_type == resource_type,
            or_(
                GovCloudImportBatch.import_mode == "full",
                GovCloudImportBatch.import_mode.is_(None),
                GovCloudImportBatch.import_mode == "",
            ),
        )
        .order_by(GovCloudImportBatch.id.desc())
        .first()
    )
    last_batch = (
        db.query(GovCloudImportBatch)
        .filter(GovCloudImportBatch.resource_type == resource_type)
        .order_by(GovCloudImportBatch.id.desc())
        .first()
    )
    if last_full is None:
        return {
            "resource_type": resource_type,
            "live_count": live,
            "excel_rows": last_batch.excel_rows if last_batch else 0,
            "unique_source_ids": last_batch.unique_source_ids if last_batch else 0,
            "consistent": live == 0,
            "last_batch": last_batch.to_dict() if last_batch else None,
            "message": (
                "尚未做过全量导入，无法按全量爬取校验一致性"
                if last_batch
                else "尚未导入过该类型，库内应为空才算一致"
            ),
        }
    consistent = live == last_full.unique_source_ids and last_full.skipped == 0
    if last_full.consistent != int(consistent) or last_full.live_count != live:
        last_full.consistent = int(consistent)
        last_full.live_count = live
        db.commit()
    return {
        "resource_type": resource_type,
        "live_count": live,
        "excel_rows": last_full.excel_rows,
        "unique_source_ids": last_full.unique_source_ids,
        "skipped": last_full.skipped,
        "consistent": consistent,
        "last_batch": last_batch.to_dict() if last_batch else None,
        "last_full": last_full.to_dict(),
        "message": (
            "与最近一次全量导入一致"
            if consistent
            else f"与最近一次全量不一致：库内 {live} 条，全量唯一 ID {last_full.unique_source_ids} 条"
        ),
    }


@router.post("/import")
async def import_govcloud_excel(
    file: UploadFile = File(..., description="爬虫产出的 xlsx / csv"),
    resource_type: str = Query("", description="cloud_host / bare_metal / e_government_network / elastic_ip"),
    mode: str = Query("full", description="导入模式：full 全量对齐 / incremental 增量 upsert 不删除"),
    request: Request = None,
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("asset_discovery", "edit")),
) -> dict:
    filename = file.filename or "upload"
    lower = filename.lower()
    is_csv = lower.endswith(".csv")
    if not (is_csv or lower.endswith(".xlsx")):
        raise HTTPException(status_code=400, detail="仅支持 .xlsx / .csv（爬虫输出格式）")
    import_mode = mode if mode in ("full", "incremental") else "full"
    raw = await file.read()
    if not raw:
        raise HTTPException(status_code=400, detail="文件为空")
    if len(raw) > MAX_FILE_SIZE:
        raise HTTPException(status_code=413, detail="文件超过 200MB")

    rtype = infer_resource_type(filename, resource_type)
    if rtype == "qingteng_host":
        raise HTTPException(
            status_code=400,
            detail="青藤主机来自 API 同步（设备管理配置青藤设备后一键拉取），不支持 Excel 导入",
        )
    model = _model(rtype)
    started = time.perf_counter()
    rows = _parse_csv(raw) if is_csv else _parse_excel(raw)
    if not rows:
        raise HTTPException(
            status_code=400,
            detail="未解析到数据行（xlsx 需 Sheet「原始数据」；csv 需首行为表头）",
        )

    skipped = 0
    seen: dict[str, dict] = {}
    duplicate = 0
    for row in rows:
        sid = _source_id(row)
        if not sid:
            skipped += 1
            continue
        if sid in seen:
            duplicate += 1
        seen[sid] = row

    unique_rows = list(seen.values())
    if not unique_rows:
        # 全部行都识别不出唯一标识：多为表头不匹配（如非爬虫口径的文件）。
        # 直接拒绝，避免产生「全量 0 条」批次并把全量模式下库内数据误删。
        raise HTTPException(
            status_code=400,
            detail=(
                f"{len(rows)} 行均无法识别唯一标识（需要 _id / instance_id / source_id 列），"
                "请确认表头与爬虫输出一致（导出的 CSV 可直接再导入）"
            ),
        )
    unique_ids = set(seen.keys())
    existing_ids = {r[0] for r in db.query(model.source_id).all()}
    inserted = len(unique_ids - existing_ids)
    updated = len(unique_ids & existing_ids)
    deleted = len(existing_ids - unique_ids) if import_mode == "full" else 0

    batch = GovCloudImportBatch(
        resource_type=rtype,
        filename=filename[:300],
        file_size=len(raw),
        excel_rows=len(rows),
        unique_source_ids=len(unique_ids),
        inserted=inserted,
        updated=updated,
        deleted=deleted,
        skipped=skipped,
        live_count=0,
        consistent=0,
        import_mode=import_mode,
        created_by=user.id,
        created_at=beijing_now(),
    )
    db.add(batch)
    db.flush()

    now = beijing_now()
    values = [_row_to_values(rtype, row, batch.id, now) for row in unique_rows]
    excel_attrs = [excel_attr(f) for f in RESOURCE_EXCEL_FIELDS[rtype]]
    update_cols = [*excel_attrs, "extra_excel", "batch_id", "updated_at"]

    for i in range(0, len(values), UPSERT_CHUNK):
        chunk = values[i:i + UPSERT_CHUNK]
        stmt = insert(model).values(chunk)
        excluded = stmt.excluded
        stmt = stmt.on_conflict_do_update(
            index_elements=["source_id"],
            set_={col: getattr(excluded, col) for col in update_cols},
        )
        db.execute(stmt)

    if import_mode == "full":
        db.query(model).filter(model.batch_id != batch.id).delete(synchronize_session=False)

    live = int(db.query(func.count(model.id)).scalar() or 0)
    consistent = live == len(unique_ids) and skipped == 0 if import_mode == "full" else skipped == 0
    notes = []
    notes.append("全量导入：已删除文件中不存在的旧记录" if import_mode == "full" else "增量导入：仅新增/更新，未删除库内既有记录")
    if not is_csv:
        notes.append("未导入 Sheet「数据统计」")
    if duplicate:
        notes.append(f"文件内重复 source_id {duplicate} 条，已按最后一条保留")
    if skipped:
        notes.append(f"缺少 _id/instance_id 跳过 {skipped} 条")
    if not consistent:
        notes.append(f"同步后库内 {live} 与唯一 ID {len(unique_ids)} 不一致")
    batch.live_count = live
    batch.consistent = int(consistent)
    batch.duration_ms = int((time.perf_counter() - started) * 1000)
    batch.note = "；".join(notes)
    db.commit()

    try:
        log_audit(
            db,
            user_id=user.id,
            username=user.username,
            action="import",
            resource_type="asset_discovery",
            resource_id=str(batch.id),
            detail=f"{rtype} {import_mode} excel={len(rows)} live={live}",
            ip_address=get_client_ip(request) if request else None,
            result="success",
        )
    except Exception as extra_exc:  # noqa: BLE001
        logger.warning("导入审计写入失败（忽略）: %s", extra_exc)
    logger.info(
        "政务云导入: mode=%s type=%s excel=%d unique=%d live=%d inserted=%d updated=%d deleted=%d consistent=%s",
        import_mode, rtype, len(rows), len(unique_ids), live, inserted, updated, deleted, consistent,
    )
    return {
        "resource_type": rtype,
        "import_mode": import_mode,
        "excel_rows": len(rows),
        "unique_source_ids": len(unique_ids),
        "inserted": inserted,
        "updated": updated,
        "deleted": deleted,
        "skipped": skipped,
        "live_count": live,
        "consistent": consistent,
        "duration_ms": batch.duration_ms,
        "note": batch.note,
        "batch": batch.to_dict(),
    }


class QingtengSyncRequest(BaseModel):
    """青藤主机同步请求：不传 device_ids 同步全部已启用的青藤设备。"""
    device_ids: list[int] = Field(default_factory=list, description="青藤设备 ID 列表（空 = 全部）")
    mode: str = Field("full", description="full 全量对齐 / incremental 增量 upsert 不删除")


@router.post("/qingteng/sync")
def sync_qingteng_hosts(
    payload: QingtengSyncRequest,
    request: Request = None,
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("asset_discovery", "edit")),
) -> dict:
    """手动触发青藤主机同步：逐台登录青藤平台，分页拉取 linux / win 主机清单入库。

    - 凭据复用设备管理（auth_type=qingteng 的 api_url/用户名/密码）；
    - 单台失败不阻断其它设备，逐台返回结果；
    - 同步口径与 Excel 导入对齐（full 删除青藤平台已不存在的记录）。
    """
    devices = get_qingteng_devices(db)
    if payload.device_ids:
        wanted = set(payload.device_ids)
        devices = [d for d in devices if d.id in wanted]
        missing = wanted - {d.id for d in devices}
        if missing:
            raise HTTPException(
                status_code=404,
                detail=f"青藤设备不存在或未启用: {sorted(missing)}，请到「设备管理」检查",
            )
    if not devices:
        raise HTTPException(
            status_code=400,
            detail="没有已启用的青藤设备：请先在「设备管理」新建设备（认证方式选择「青藤签名」），"
                   "配置 API 地址与 Console 用户名/密码后启用",
        )

    import_mode = payload.mode if payload.mode in ("full", "incremental") else "full"
    results = []
    for dev in devices:
        try:
            res = sync_device_hosts(db, dev, user_id=user.id, mode=import_mode)
            results.append({"success": True, **res})
        except Exception as exc:  # noqa: BLE001
            logger.warning("青藤同步失败: device=%s(%d) err=%s", dev.name, dev.id, exc)
            results.append({
                "success": False,
                "device_id": dev.id,
                "device_name": dev.name,
                "import_mode": import_mode,
                "error": str(exc),
            })

    ok_count = sum(1 for r in results if r.get("success"))
    try:
        log_audit(
            db,
            user_id=user.id,
            username=user.username,
            action="sync",
            resource_type="asset_discovery",
            detail=(
                f"qingteng {import_mode} devices={len(devices)} ok={ok_count} "
                f"failed={len(results) - ok_count}"
            ),
            ip_address=get_client_ip(request) if request else None,
            result="success" if ok_count == len(results) else "partial",
        )
    except Exception as exc:  # noqa: BLE001
        logger.warning("同步审计写入失败（忽略）: %s", exc)

    return {
        "success": ok_count == len(results),
        "mode": import_mode,
        "device_count": len(devices),
        "ok_count": ok_count,
        "results": results,
    }
