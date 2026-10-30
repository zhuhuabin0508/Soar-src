"""文档生成（generate_document 工具宿主能力）单测。

覆盖：
- Markdown 块解析纯函数（标题/段落/有序无序列表/表格/代码块）；
- 文件名清洗（非法字符剥离、扩展名规范、空默认值、超长截断）；
- 参数校验（空内容 / 非法格式 / 超长 / options 类型错误均返回明确 error，不抛异常）；
- 三种格式（md / docx / xlsx）的生成与文件内容正确性
  （monkeypatch 上传目录与 SessionLocal，不依赖真实数据库）。

依赖说明：docx / openpyxl 为后端运行依赖（read_uploaded_file 同源），
宿主机执行 pytest 需确保已安装（backend/requirements）。
"""
import pytest

from app.core.tool_runner import (
    _coerce_cell,
    _parse_markdown_blocks,
    _sanitize_doc_name,
    save_document_as_agent_file,
)

SAMPLE = """# 概况
本期共发现 3 个风险。

## 风险清单
| IP | 等级 | 端口 |
| --- | --- | --- |
| 1.1.1.1 | 高 | 8080 |
| 2.2.2.2 | 低 | 443 |

- 第一项
- 第二项

1. 步骤一
2. 步骤二

**重点**：请尽快处置。

```python
print("hello")
```
"""


# ---------------------------------------------------------------------------
# Markdown 块解析（纯函数）
# ---------------------------------------------------------------------------

def test_parse_markdown_blocks_full_document():
    blocks = _parse_markdown_blocks(SAMPLE)
    kinds = [b["kind"] for b in blocks]
    assert kinds == [
        "heading", "paragraph", "heading", "table",
        "list", "list", "paragraph", "code",
    ]
    assert blocks[0] == {"kind": "heading", "level": 1, "text": "概况"}
    assert blocks[1]["text"] == "本期共发现 3 个风险。"
    assert blocks[2] == {"kind": "heading", "level": 2, "text": "风险清单"}
    # 表格：分隔行剔除，首行为表头
    assert blocks[3]["rows"] == [
        ["IP", "等级", "端口"],
        ["1.1.1.1", "高", "8080"],
        ["2.2.2.2", "低", "443"],
    ]
    assert blocks[4] == {"kind": "list", "ordered": False, "items": ["第一项", "第二项"]}
    assert blocks[5] == {"kind": "list", "ordered": True, "items": ["步骤一", "步骤二"]}
    # 行内 **粗体** 保留原样（由 docx 生成器拆 run）
    assert blocks[6]["text"] == "**重点**：请尽快处置。"
    assert blocks[7] == {"kind": "code", "text": 'print("hello")'}


def test_parse_markdown_blocks_empty_and_blank():
    assert _parse_markdown_blocks("") == []
    assert _parse_markdown_blocks("  \n\n \n") == []


def test_parse_markdown_blocks_single_line_table_without_separator():
    """只有一个 | 行（无分隔行）时不是表格，按段落处理（连续普通行合并），避免误判。"""
    blocks = _parse_markdown_blocks("|a|b|\n正文")
    assert blocks[0]["kind"] == "paragraph"
    assert blocks[0]["text"] == "|a|b|\n正文"


# ---------------------------------------------------------------------------
# 文件名清洗
# ---------------------------------------------------------------------------

def test_sanitize_doc_name_strips_unsafe_and_normalizes_ext():
    assert _sanitize_doc_name("报告<>:v1.txt", "docx") == "报告v1.docx"
    # 路径分隔符剔除 + 首尾点剥离（无路径穿越残留）
    assert _sanitize_doc_name("../evil/path", "md") == "evilpath.md"
    assert _sanitize_doc_name("  数据导出  ", "xlsx") == "数据导出.xlsx"
    # 空名给默认（带时间戳前缀）
    default = _sanitize_doc_name("   ", "docx")
    assert default.startswith("agent文档_") and default.endswith(".docx")
    # 超长截断到 80 字符 + 扩展名
    long_name = "字" * 100
    assert len(_sanitize_doc_name(long_name, "md")) <= 83


# ---------------------------------------------------------------------------
# 参数校验（不触发文件生成）
# ---------------------------------------------------------------------------

def test_save_document_rejects_empty_content():
    for bad in ("", "   ", None):
        result = save_document_as_agent_file(bad, "报告")
        assert "error" in result
        assert "不能为空" in result["error"]


def test_save_document_rejects_invalid_format():
    result = save_document_as_agent_file("正文", "报告", fmt="pptx")
    assert "error" in result
    assert "仅支持 docx / xlsx / md" in result["error"]


def test_save_document_rejects_oversized_content():
    result = save_document_as_agent_file("x" * 400_001, "报告")
    assert "error" in result
    assert "过长" in result["error"]


def test_save_document_rejects_bad_options_type():
    result = save_document_as_agent_file("正文", "报告", options=["not", "a", "dict"])
    assert "error" in result
    assert "options" in result["error"]


# ---------------------------------------------------------------------------
# 三种格式生成（monkeypatch 上传目录 + SessionLocal）
# ---------------------------------------------------------------------------

class _FakeSession:
    """替代 SessionLocal：捕获 AgentFile 记录并赋予固定 id。"""

    rec = None

    def add(self, obj):
        self.rec = obj

    def commit(self):
        pass

    def refresh(self, obj):
        obj.id = 4242

    def rollback(self):
        pass

    def close(self):
        pass


@pytest.fixture()
def fake_file_store(tmp_path, monkeypatch):
    """把上传目录指到 tmp_path、SessionLocal 换成 fake，返回 (dir, fake)。"""
    fake = _FakeSession()
    monkeypatch.setattr("app.core.tool_runner._agent_upload_dir", lambda: str(tmp_path))
    monkeypatch.setattr("app.database.SessionLocal", lambda: fake)
    return tmp_path, fake


def test_save_document_markdown(fake_file_store):
    tmp_path, fake = fake_file_store
    result = save_document_as_agent_file(
        SAMPLE, "测试报告.md", fmt="md", options={"title": "周报"}, agent_id=7, user_id=3
    )
    assert result.get("ok") is True
    assert result["file_name"] == "测试报告.md"
    assert result["file_id"] == 4242
    assert result["download_url"] == f"/api/v1/agent-files/4242/download"
    # 归属隔离：记录写入当前智能体与用户
    assert fake.rec.agent_id == 7 and fake.rec.created_by == 3
    # title 注入顶部大标题，正文原样保留
    stored = tmp_path / fake.rec.stored_name
    assert stored.exists()
    content = stored.read_text(encoding="utf-8")
    assert content.startswith("# 周报\n\n")
    assert "本期共发现 3 个风险。" in content


def test_save_document_docx(fake_file_store):
    import docx as docx_lib

    tmp_path, fake = fake_file_store
    result = save_document_as_agent_file(SAMPLE, "渗透报告", fmt="docx")
    assert result.get("ok") is True
    assert result["file_type"] == "docx"
    path = tmp_path / fake.rec.stored_name
    assert path.exists() and path.stat().st_size > 0

    doc = docx_lib.Document(str(path))
    # 标题块 + 表格 + 段落齐备，无内容丢失
    texts = [p.text for p in doc.paragraphs]
    assert "概况" in texts
    assert "本期共发现 3 个风险。" in texts
    assert "1. 步骤一" in texts and "• 第一项" in texts
    assert 'print("hello")' in texts
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["IP", "等级", "端口"]
    assert table.cell(1, 0).text == "1.1.1.1"
    # 行内粗体拆 run：重点为粗体 run
    bold_runs = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
    assert "重点" in bold_runs
    # 中文字体（eastAsia）已设置到 Normal 样式
    from docx.oxml.ns import qn
    normal = doc.styles["Normal"]
    assert normal.font.name == "微软雅黑"
    assert normal.element.rPr.rFonts.get(qn("w:eastAsia")) == "微软雅黑"


def test_save_document_xlsx(fake_file_store):
    from openpyxl import load_workbook

    tmp_path, fake = fake_file_store
    result = save_document_as_agent_file(SAMPLE, "风险清单", fmt="xlsx")
    assert result.get("ok") is True
    assert result["file_type"] == "xlsx"
    path = tmp_path / fake.rec.stored_name

    wb = load_workbook(str(path))
    # 表格独立成表（命名取就近标题），文本进「文档」表，两者互不覆盖
    assert "风险清单" in wb.sheetnames
    assert "文档" in wb.sheetnames
    ws = wb["风险清单"]
    assert ws.cell(1, 1).value == "IP"
    assert ws.cell(2, 1).value == "1.1.1.1"
    assert ws.cell(2, 3).value == 8080  # 纯数字短串转数值
    assert ws.freeze_panes == "A2"
    assert ws.cell(1, 1).font.bold is True
    text_ws = wb["文档"]
    texts = [text_ws.cell(r, 1).value for r in range(1, 6)]
    assert "概况" in texts and "本期共发现 3 个风险。" in texts


def test_save_document_xlsx_table_only_removes_empty_text_sheet(fake_file_store):
    """纯表格内容：空的「文档」表被移除，工作簿仍有表格表。"""
    from openpyxl import load_workbook

    tmp_path, fake = fake_file_store
    content = "| A | B |\n| --- | --- |\n| 1 | 2 |\n"
    result = save_document_as_agent_file(content, "纯表格", fmt="xlsx")
    assert result.get("ok") is True
    wb = load_workbook(str(tmp_path / fake.rec.stored_name))
    assert "文档" not in wb.sheetnames
    assert wb.sheetnames == ["表格1"]
    assert wb["表格1"].cell(2, 1).value == 1


def test_save_document_format_alias_and_default(fake_file_store):
    """markdown/word 别名与默认 docx。"""
    tmp_path, fake = fake_file_store
    r1 = save_document_as_agent_file("正文一", "别名", fmt="markdown")
    assert r1["file_type"] == "md"
    r2 = save_document_as_agent_file("正文二", "默认")
    assert r2["file_type"] == "docx"


def test_coerce_cell():
    assert _coerce_cell("8080") == 8080
    assert _coerce_cell("3.14") == 3.14
    assert _coerce_cell("0") == 0
    assert _coerce_cell("1.1.1.1") == "1.1.1.1"  # IP 保留字符串
    assert _coerce_cell("007") == "007"  # 前导零保留字符串
    assert _coerce_cell("1234567890123456") == "1234567890123456"  # 超长数字保留字符串
    assert _coerce_cell("高") == "高"
